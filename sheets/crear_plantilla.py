import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.worksheet.datavalidation import DataValidation

MUESTRA = len(sys.argv) > 2 and sys.argv[2] == 'muestra'
F = 'Arial'
AZUL = '1F3A5F'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor=AZUL)
N = Font(name=F); B = Font(name=F, bold=True); T = Font(name=F, bold=True, size=16, color=AZUL)
GRIS = Font(name=F, italic=True, color='7F7F7F', size=9)
BOT = PatternFill('solid', fgColor='F2F2F2')
SOLES = '"S/ "#,##0.00;"-S/ "#,##0.00;"-"'
FECHAHORA = 'dd/mm/yyyy hh:mm'; FECHA = 'dd/mm/yyyy'
CATS = ['Comida', 'Transporte', 'Ocio', 'Compras', 'Salud', 'Educación', 'Servicios', 'Hogar', 'Otros']
COLORES = ['2A78D6', 'EB6834', '1BAF7A', 'EDA100', 'E87BA4', '008300', '4A3AA7', 'E34948', '8A8986']
FILAS = 3000

wb = Workbook()

def encabezado(ws, fila, cols, anchos):
    for i, (c, w) in enumerate(zip(cols, anchos), start=1):
        cell = ws.cell(row=fila, column=i, value=c); cell.font = H; cell.fill = HF
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[fila].height = 20

def hoja_datos(nombre, cols, anchos, formatos, nota=True, desde_col=1):
    ws = wb.create_sheet(nombre)
    encabezado(ws, 1, cols, anchos)
    ws.freeze_panes = 'A2'
    for r in range(2, FILAS + 2 if nombre == 'Movimientos' else 402):
        for c, fmt in formatos.items():
            ws.cell(row=r, column=c).number_format = fmt
    return ws

# ---------------- Resumen ----------------
ws = wb.active; ws.title = 'Resumen'
ws['A1'] = 'Mis finanzas'; ws['A1'].font = T
ws['A2'] = 'Última actualización:'; ws['A2'].font = GRIS
ws['B2'].font = GRIS
ws['A3'] = 'La llena tu bot de Telegram automáticamente. Las pestañas de datos no se editan a mano; las correcciones se hacen por el bot (/borrar, /cat…).'
ws['A3'].font = GRIS
ws['A5'] = 'Inicio del mes'; ws['B5'] = '=DATE(YEAR(TODAY()),MONTH(TODAY()),1)'; ws['B5'].number_format = FECHA
ws['A6'] = 'Inicio de la semana (lunes)'; ws['B6'] = '=TODAY()-WEEKDAY(TODAY(),3)'; ws['B6'].number_format = FECHA
for c in ('A5', 'A6', 'B5', 'B6'): ws[c].font = GRIS
kpis = [
    ('Gastos del mes', '=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$A:$A,">="&$B$5)'),
    ('Ingresos del mes', '=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Ingreso",Movimientos!$A:$A,">="&$B$5)'),
    ('Balance del mes', '=B9-B8'),
    ('Gastos de esta semana', '=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$A:$A,">="&$B$6)'),
    ('Promedio diario (mes)', '=IFERROR(B8/DAY(TODAY()),0)'),
    ('Gastos del mes anterior', '=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$A:$A,">="&EDATE($B$5,-1),Movimientos!$A:$A,"<"&$B$5)'),
    ('Te deben', '=SUMIF(Deudas!$B:$B,">0")'),
    ('Debes', '=-SUMIF(Deudas!$B:$B,"<0")'),
    ('Ahorrado en metas', '=SUM(Metas!$C:$C)'),
]
ws['A7'] = 'Indicador'; ws['B7'] = 'Valor'
for c in ('A7', 'B7'): ws[c].font = H; ws[c].fill = HF
for i, (k, f) in enumerate(kpis, start=8):
    ws.cell(row=i, column=1, value=k).font = N
    c = ws.cell(row=i, column=2, value=f); c.font = B; c.number_format = SOLES
    if i % 2: 
        ws.cell(row=i, column=1).fill = BOT; c.fill = BOT
ws.column_dimensions['A'].width = 28; ws.column_dimensions['B'].width = 18

# ---------------- Datos (los llena el bot) ----------------
mov = hoja_datos('Movimientos', ['Fecha', 'Tipo', 'Monto (S/)', 'Categoría', 'Descripción', 'Método', 'Semana (lunes)'],
                 [17, 10, 13, 14, 34, 12, 15], {1: FECHAHORA, 3: SOLES, 7: FECHA})
pres = hoja_datos('Presupuestos', ['Categoría', 'Presupuesto mensual (S/)'], [16, 24], {2: SOLES})
deu = wb.create_sheet('Deudas')
encabezado(deu, 1, ['Persona', 'Saldo (S/)', 'Estado', '', 'Fecha', 'Persona', 'Monto (S/)', 'Descripción'], [16, 13, 14, 3, 17, 16, 13, 30])
deu['D1'].fill = PatternFill(fill_type=None)
deu.freeze_panes = 'A2'
for r in range(2, 402):
    deu.cell(row=r, column=2).number_format = SOLES; deu.cell(row=r, column=5).number_format = FECHAHORA; deu.cell(row=r, column=7).number_format = SOLES
metas = hoja_datos('Metas', ['Meta', 'Objetivo (S/)', 'Ahorrado (S/)', 'Fecha límite', 'Progreso', 'Falta (S/)', 'Ahorrar por mes (S/)'],
                   [20, 14, 14, 14, 11, 13, 20], {2: SOLES, 3: SOLES, 4: FECHA, 5: '0%', 6: SOLES, 7: SOLES})
for r in range(2, 402):
    metas[f'E{r}'] = f'=IF(A{r}="","",IFERROR(C{r}/B{r},0))'
    metas[f'F{r}'] = f'=IF(A{r}="","",MAX(B{r}-C{r},0))'
    metas[f'G{r}'] = f'=IF(OR(A{r}="",D{r}=""),"",IFERROR(F{r}/(DATEDIF(TODAY(),D{r},"m")+1),F{r}))'
tar = hoja_datos('Tareas', ['Tarea', 'Prioridad', 'Vence', 'Estado', 'Hecha el', 'Creada'], [40, 11, 13, 11, 17, 17],
                 {3: FECHA, 5: FECHAHORA, 6: FECHAHORA})
rec = hoja_datos('Recurrentes', ['Pago', 'Monto (S/)', 'Categoría', 'Día del mes', 'Activo'], [22, 13, 14, 12, 9], {2: SOLES})
rdo = hoja_datos('Recordatorios', ['Recordatorio', 'Próxima vez', 'Se repite'], [44, 17, 14], {2: FECHAHORA})

# ---------------- Mensual ----------------
men = wb.create_sheet('Mensual', 1)
encabezado(men, 1, ['Mes', 'Ingresos (S/)', 'Gastos (S/)', 'Balance (S/)', '% gastado'], [12, 15, 15, 15, 12])
for i in range(12):
    r = i + 2
    men[f'A{r}'] = f'=EDATE(Resumen!$B$5,{i - 11})'; men[f'A{r}'].number_format = 'mmm yyyy'
    men[f'B{r}'] = f'=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Ingreso",Movimientos!$A:$A,">="&A{r},Movimientos!$A:$A,"<"&EDATE(A{r},1))'
    men[f'C{r}'] = f'=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$A:$A,">="&A{r},Movimientos!$A:$A,"<"&EDATE(A{r},1))'
    men[f'D{r}'] = f'=B{r}-C{r}'
    men[f'E{r}'] = f'=IF(B{r}>0,C{r}/B{r},"")'
    for c in 'BCD': men[f'{c}{r}'].number_format = SOLES
    men[f'E{r}'].number_format = '0%'
men['A14'] = 'Total 12 meses'; men['A14'].font = B
for c in 'BCD':
    men[f'{c}14'] = f'=SUM({c}2:{c}13)'; men[f'{c}14'].font = B; men[f'{c}14'].number_format = SOLES
men.freeze_panes = 'A2'

# ---------------- Semanal ----------------
sem = wb.create_sheet('Semanal', 2)
encabezado(sem, 1, ['Semana (lunes)', 'Ingresos (S/)', 'Gastos (S/)', 'Balance (S/)'], [16, 15, 15, 15])
for i in range(12):
    r = i + 2
    sem[f'A{r}'] = f'=Resumen!$B$6-7*{11 - i}'; sem[f'A{r}'].number_format = FECHA
    sem[f'B{r}'] = f'=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Ingreso",Movimientos!$A:$A,">="&A{r},Movimientos!$A:$A,"<"&(A{r}+7))'
    sem[f'C{r}'] = f'=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$A:$A,">="&A{r},Movimientos!$A:$A,"<"&(A{r}+7))'
    sem[f'D{r}'] = f'=B{r}-C{r}'
    for c in 'BCD': sem[f'{c}{r}'].number_format = SOLES
sem['A14'] = 'Promedio semanal'; sem['A14'].font = B
sem['C14'] = '=AVERAGE(C2:C13)'; sem['C14'].font = B; sem['C14'].number_format = SOLES
sem.freeze_panes = 'A2'

# ---------------- Categorías (mes actual) ----------------
cat = wb.create_sheet('Categorías', 3)
encabezado(cat, 1, ['Categoría', 'Gasto del mes (S/)', 'Presupuesto (S/)', '% usado', 'Te queda (S/)', 'Mes anterior (S/)', 'vs mes anterior'],
           [14, 18, 16, 10, 15, 17, 15])
for i, nombre in enumerate(CATS):
    r = i + 2
    cat[f'A{r}'] = nombre
    cat[f'B{r}'] = f'=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$D:$D,A{r},Movimientos!$A:$A,">="&Resumen!$B$5)'
    cat[f'C{r}'] = f'=IFERROR(INDEX(Presupuestos!$B:$B,MATCH(A{r},Presupuestos!$A:$A,0)),0)'
    cat[f'D{r}'] = f'=IF(C{r}>0,B{r}/C{r},"")'
    cat[f'E{r}'] = f'=IF(C{r}>0,C{r}-B{r},"")'
    cat[f'F{r}'] = f'=SUMIFS(Movimientos!$C:$C,Movimientos!$B:$B,"Gasto",Movimientos!$D:$D,A{r},Movimientos!$A:$A,">="&EDATE(Resumen!$B$5,-1),Movimientos!$A:$A,"<"&Resumen!$B$5)'
    cat[f'G{r}'] = f'=IF(F{r}>0,B{r}/F{r}-1,"")'
    for c in 'BCEF': cat[f'{c}{r}'].number_format = SOLES
    cat[f'D{r}'].number_format = '0%'; cat[f'G{r}'].number_format = '+0%;-0%;0%'
cat['A11'] = 'Total'; cat['A11'].font = B
for c in 'BCF':
    cat[f'{c}11'] = f'=SUM({c}2:{c}10)'; cat[f'{c}11'].font = B; cat[f'{c}11'].number_format = SOLES
cat.freeze_panes = 'A2'

# ---------------- Gráficos en Resumen ----------------
pie = PieChart(); pie.title = 'Gastos del mes por categoría'
pie.add_data(Reference(cat, min_col=2, min_row=1, max_row=10), titles_from_data=True)
pie.set_categories(Reference(cat, min_col=1, min_row=2, max_row=10))
for i, col in enumerate(COLORES):
    pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = col; pie.series[0].dPt.append(pt)
pie.height = 9; pie.width = 15
ws.add_chart(pie, 'D1')

bar = BarChart(); bar.type = 'col'; bar.title = 'Ingresos vs gastos (últimos 12 meses)'
bar.add_data(Reference(men, min_col=2, max_col=3, min_row=1, max_row=13), titles_from_data=True)
bar.set_categories(Reference(men, min_col=1, min_row=2, max_row=13))
bar.series[0].graphicalProperties.solidFill = '2A78D6'; bar.series[1].graphicalProperties.solidFill = 'EB6834'
bar.y_axis.title = 'S/'; bar.height = 9; bar.width = 15
ws.add_chart(bar, 'D19')

bar2 = BarChart(); bar2.type = 'col'; bar2.title = 'Gastos por semana (últimas 12)'
bar2.add_data(Reference(sem, min_col=3, min_row=1, max_row=13), titles_from_data=True)
bar2.set_categories(Reference(sem, min_col=1, min_row=2, max_row=13))
bar2.series[0].graphicalProperties.solidFill = '2A78D6'; bar2.legend = None; bar2.height = 9; bar2.width = 15
ws.add_chart(bar2, 'D37')

for s in wb.worksheets:
    for row in s.iter_rows(min_row=1, max_row=min(s.max_row, 20)):
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None,
                              italic=c.font.italic if c.font else False, size=c.font.size if c.font else 11)

# ---------------- Datos de muestra (solo para probar fórmulas) ----------------
if MUESTRA:
    import datetime as dt
    hoy = dt.date.today()
    base = dt.datetime(1899, 12, 30)
    def serial(d):
        return (d - base).total_seconds() / 86400
    ahora = dt.datetime.now()
    filas = [
        (ahora - dt.timedelta(hours=1), 'Gasto', 15, 'Comida', 'almuerzo', 'manual'),
        (ahora - dt.timedelta(days=1), 'Gasto', 8.5, 'Transporte', 'uber', 'correo'),
        (ahora - dt.timedelta(hours=2), 'Ingreso', 1500, '', 'sueldo', ''),
        (ahora.replace(day=1, hour=9) - dt.timedelta(days=5), 'Gasto', 120, 'Compras', 'zapatillas', 'manual'),
        (ahora.replace(day=1, hour=9) - dt.timedelta(days=5), 'Ingreso', 1400, '', 'sueldo', ''),
    ]
    for r, f in enumerate(filas, start=2):
        lunes = (f[0].date() - dt.timedelta(days=f[0].weekday()))
        vals = [serial(f[0]), f[1], f[2], f[3], f[4], f[5], serial(dt.datetime.combine(lunes, dt.time()))]
        for c, v in enumerate(vals, start=1): mov.cell(row=r, column=c, value=v)
    pres['A2'] = 'Comida'; pres['B2'] = 40
    deu['A2'] = 'Juan'; deu['B2'] = 30; deu['C2'] = 'te debe'; deu['A3'] = 'Ana'; deu['B3'] = -12; deu['C3'] = 'le debes'
    metas['A2'] = 'laptop'; metas['B2'] = 2000; metas['C2'] = 250; metas['D2'] = serial(dt.datetime(hoy.year, 12, 31))

wb.save(sys.argv[1])
print('guardado', sys.argv[1])
